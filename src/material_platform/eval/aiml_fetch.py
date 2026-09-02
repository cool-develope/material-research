from __future__ import annotations

import json
import tarfile
from io import BytesIO
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen
from zipfile import ZIP_DEFLATED, ZipFile

from material_platform.eval.aiml_corpus import (
    AIML_CACHE,
    AIML_GENERATED_SUITE,
    AIML_ZIP,
    SOURCES,
    TRANSFORMERS_ROOT,
    RemoteFile,
    make_suite,
)

USER_AGENT = "MaterialPlatformEval/0.1 (research fixture builder)"
ALPACA_ROWS = 200
TRANSFORMERS_MAX_FILES = 4_500
DEFAULT_GITHUB_MAX_FILES = 2_000
_SKIP_PARTS = frozenset(
    {".git", ".github", "docs", "notebooks", "docker", "node_modules", "dist"}
)
_SKIP_SUFFIXES = frozenset(
    {
        ".png",
        ".jpg",
        ".jpeg",
        ".gif",
        ".webp",
        ".svg",
        ".bin",
        ".pt",
        ".onnx",
        ".safetensors",
        ".pkl",
        ".h5",
        ".npz",
        ".woff",
        ".ttf",
        ".mp4",
        ".wav",
        ".lock",
        ".wasm",
        ".map",
    }
)
_IDENTITY_FILES = frozenset(
    {
        "readme.md",
        "readme.rst",
        "license",
        "license.txt",
        "licence",
        "pyproject.toml",
        "setup.py",
        "setup.cfg",
        "package.json",
        "go.mod",
        "cargo.toml",
    }
)
_NPM_KEEP_SUFFIXES = frozenset({".json", ".md", ".js", ".ts", ".mjs", ".cjs"})


def build_aiml_zip(
    *,
    cache: Path = AIML_CACHE,
    zip_path: Path = AIML_ZIP,
    suite_path: Path = AIML_GENERATED_SUITE,
) -> Path:
    tree = cache / "tree"
    _reset_tree(tree)
    for source in SOURCES:
        _materialize(source, cache, tree)
    _write_readme(tree)
    zip_path.parent.mkdir(parents=True, exist_ok=True)
    _zip_tree(tree, zip_path)
    suite_path.write_text(
        json.dumps(make_suite().model_dump(mode="json"), indent=2) + "\n",
        encoding="utf-8",
    )
    return zip_path


_CHUNK = 64 * 1024


def download_bytes(url: str, *, cache_file: Path | None = None) -> bytes:
    if (
        cache_file is not None
        and cache_file.is_file()
        and cache_file.stat().st_size > 0
    ):
        return cache_file.read_bytes()
    if cache_file is None:
        return _http_get(url)
    cache_file.parent.mkdir(parents=True, exist_ok=True)
    tmp = cache_file.with_name(cache_file.name + ".part")
    _http_get(url, dest=tmp)
    tmp.replace(cache_file)
    return cache_file.read_bytes()


def alpaca_to_csv(payload: object, *, rows: int = ALPACA_ROWS) -> str:
    if not isinstance(payload, list):
        raise ValueError("alpaca json must be a list of examples")
    lines = ["instruction,input,output"]
    taken = 0
    for item in payload:
        if not isinstance(item, dict):
            continue
        instruction = str(item.get("instruction") or "")
        if not instruction:
            continue
        lines.append(
            ",".join(
                (
                    _csv_cell(instruction),
                    _csv_cell(str(item.get("input") or "")),
                    _csv_cell(str(item.get("output") or "")),
                )
            )
        )
        taken += 1
        if taken >= rows:
            break
    if taken == 0:
        raise ValueError("alpaca json had no instruction rows")
    return "\n".join(lines) + "\n"


def markdown_to_docx(text: str, *, title: str) -> bytes:
    from docx import Document

    document = Document()
    document.core_properties.title = title
    document.core_properties.author = "Hugging Face TRL"
    for raw in text.splitlines():
        line = raw.rstrip()
        if line.startswith("#"):
            hashes, _, heading = line.partition(" ")
            level = min(len(hashes), 3)
            document.add_heading(heading.strip() or line, level=level)
        elif line:
            document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def select_github_members(
    names: tuple[str, ...],
    *,
    max_files: int,
    prefer: tuple[str, ...] = ("src/", "tests/", "examples/"),
) -> tuple[str, ...]:
    relative: list[tuple[str, str]] = []
    for name in names:
        if name.endswith("/"):
            continue
        rel = _strip_first(name)
        if not rel or _skip_member(rel):
            continue
        relative.append((name, rel))
    identity = [
        (name, rel)
        for name, rel in relative
        if Path(rel).name.lower() in _IDENTITY_FILES and "/" not in rel
    ]
    chosen: list[str] = []
    seen: set[str] = set()
    for name, rel in identity:
        if len(chosen) >= max_files:
            break
        chosen.append(name)
        seen.add(rel)
    for prefix in prefer:
        for name, rel in relative:
            if rel in seen or not rel.startswith(prefix):
                continue
            if len(chosen) >= max_files:
                return tuple(chosen)
            chosen.append(name)
            seen.add(rel)
    for name, rel in relative:
        if rel in seen:
            continue
        if len(chosen) >= max_files:
            break
        chosen.append(name)
        seen.add(rel)
    return tuple(chosen)


def _materialize(source: RemoteFile, cache: Path, tree: Path) -> None:
    dest = tree / source.dest
    if source.kind == "pdf":
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(_cached(source, cache, suffix=".pdf"))
        return
    if source.kind == "md":
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(_cached(source, cache, suffix=".md"))
        return
    if source.kind == "md_as_docx":
        dest.parent.mkdir(parents=True, exist_ok=True)
        text = _cached(source, cache, suffix=".md").decode("utf-8", errors="replace")
        dest.write_bytes(markdown_to_docx(text, title=dest.stem))
        return
    if source.kind == "alpaca_json":
        dest.parent.mkdir(parents=True, exist_ok=True)
        payload = json.loads(_cached(source, cache, suffix=".json"))
        dest.write_text(alpaca_to_csv(payload), encoding="utf-8")
        return
    if source.kind == "github_zip":
        archive = _cached(source, cache, suffix=".zip")
        max_files = (
            TRANSFORMERS_MAX_FILES
            if source.dest == TRANSFORMERS_ROOT
            else DEFAULT_GITHUB_MAX_FILES
        )
        _extract_github(archive, dest, max_files=max_files)
        return
    if source.kind == "pypi_wheel":
        dest.mkdir(parents=True, exist_ok=True)
        _fetch_wheel(source, cache, dest)
        return
    if source.kind == "jar":
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(_cached(source, cache, suffix=".jar"))
        return
    if source.kind == "npm_tgz":
        dest.mkdir(parents=True, exist_ok=True)
        _extract_npm(_cached(source, cache, suffix=".tgz"), dest)
        return
    raise ValueError(f"unknown source kind {source.kind}")


def _cached(source: RemoteFile, cache: Path, *, suffix: str) -> bytes:
    stem = source.dest.rstrip("/").replace("/", "__") or "download"
    path = cache / "downloads" / f"{stem}{suffix}"
    last_error: Exception | None = None
    if path.is_file() and path.stat().st_size > 0:
        return path.read_bytes()
    for url in (source.url, *source.fallbacks):
        for _attempt in range(3):
            try:
                print(f"downloading {url}", flush=True)
                return download_bytes(url, cache_file=path)
            except (HTTPError, URLError, TimeoutError, OSError) as exc:
                last_error = exc
                print(f"  failed: {exc}", flush=True)
                part = path.with_name(path.name + ".part")
                if part.exists():
                    part.unlink()
                if isinstance(exc, HTTPError):
                    break
    raise RuntimeError(f"could not download {source.dest}: {last_error}")


def _fetch_wheel(source: RemoteFile, cache: Path, dest_dir: Path) -> None:
    meta_path = cache / "downloads" / f"{source.url.split('/')[-2]}-pypi.json"
    raw = download_bytes(source.url, cache_file=meta_path)
    payload = json.loads(raw)
    urls = payload.get("urls")
    if not isinstance(urls, list):
        raise ValueError(f"no urls in {source.url}")
    wheel = _pick_wheel(urls)
    filename = str(wheel["filename"])
    cached = cache / "downloads" / filename
    if not cached.is_file():
        print(f"downloading {wheel['url']}", flush=True)
        download_bytes(str(wheel["url"]), cache_file=cached)
    target = dest_dir / filename
    target.write_bytes(cached.read_bytes())


def _pick_wheel(urls: list[object]) -> dict[str, object]:
    wheels = [
        item
        for item in urls
        if isinstance(item, dict)
        and item.get("packagetype") == "bdist_wheel"
        and isinstance(item.get("url"), str)
        and isinstance(item.get("filename"), str)
    ]
    if not wheels:
        raise ValueError("no wheel on PyPI")
    any_wheel = [
        item
        for item in wheels
        if "py3-none-any" in str(item["filename"])
        or "py2.py3-none-any" in str(item["filename"])
    ]
    return any_wheel[0] if any_wheel else wheels[0]


def _extract_github(data: bytes, dest: Path, *, max_files: int) -> None:
    dest.mkdir(parents=True, exist_ok=True)
    with ZipFile(BytesIO(data)) as archive:
        names = select_github_members(tuple(archive.namelist()), max_files=max_files)
        for name in names:
            relative = _strip_first(name)
            target = _safe_join(dest, relative)
            target.parent.mkdir(parents=True, exist_ok=True)
            with archive.open(name) as handle:
                target.write_bytes(handle.read())


def _extract_npm(data: bytes, dest: Path) -> None:
    dest.mkdir(parents=True, exist_ok=True)
    with tarfile.open(fileobj=BytesIO(data), mode="r:gz") as archive:
        for member in archive.getmembers():
            if not member.isfile():
                continue
            relative = _strip_first(member.name)
            if not relative or _skip_member(relative):
                continue
            suffix = Path(relative).suffix.lower()
            if suffix not in _NPM_KEEP_SUFFIXES and Path(relative).name != "LICENSE":
                continue
            target = _safe_join(dest, relative)
            extracted = archive.extractfile(member)
            if extracted is None:
                continue
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(extracted.read())


def _zip_tree(tree: Path, zip_path: Path) -> None:
    with ZipFile(zip_path, "w", compression=ZIP_DEFLATED) as archive:
        for file in sorted(tree.rglob("*")):
            if file.is_file():
                archive.write(file, file.relative_to(tree).as_posix())


def _write_readme(tree: Path) -> None:
    lines = [
        "# AI/ML eval corpus",
        "",
        "Downloaded public sources for local ingest. Not redistributed in git.",
        "SLP3 draft: class/print use per the authors; all rights reserved.",
        "",
        "Sources:",
    ]
    for source in SOURCES:
        lines.append(f"- `{source.dest}` ← {source.url}")
    (tree / "README.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def _reset_tree(tree: Path) -> None:
    if tree.exists():
        for item in sorted(tree.rglob("*"), reverse=True):
            if item.is_file() or item.is_symlink():
                item.unlink()
            elif item.is_dir():
                item.rmdir()
        tree.rmdir()
    tree.mkdir(parents=True)


def _http_get(url: str, *, dest: Path | None = None, timeout: int = 60) -> bytes:
    request = Request(
        url,
        headers={
            "User-Agent": USER_AGENT,
            "Accept": "*/*",
        },
    )
    with urlopen(request, timeout=timeout) as response:
        if dest is None:
            chunks: list[bytes] = []
            while True:
                chunk = response.read(_CHUNK)
                if not chunk:
                    break
                if not isinstance(chunk, bytes):
                    raise TypeError("expected bytes from download")
                chunks.append(chunk)
            return b"".join(chunks)
        dest.parent.mkdir(parents=True, exist_ok=True)
        with dest.open("wb") as handle:
            while True:
                chunk = response.read(_CHUNK)
                if not chunk:
                    break
                handle.write(chunk)
        return dest.read_bytes()


def _strip_first(name: str) -> str:
    parts = name.replace("\\", "/").split("/")
    if len(parts) <= 1:
        return ""
    return "/".join(parts[1:])


def _skip_member(relative: str) -> bool:
    parts = Path(relative).parts
    if any(part in _SKIP_PARTS for part in parts):
        return True
    suffix = Path(relative).suffix.lower()
    return suffix in _SKIP_SUFFIXES


def _safe_join(root: Path, relative: str) -> Path:
    root = root.resolve()
    target = (root / relative).resolve()
    if not target.is_relative_to(root):
        raise ValueError(f"unsafe archive path: {relative}")
    return target


def _csv_cell(value: str) -> str:
    escaped = value.replace('"', '""')
    return f'"{escaped}"'
