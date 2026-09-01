from io import BytesIO

from docx import Document


def build_docx(
    *,
    paragraphs: tuple[str, ...] = (),
    sections: tuple[tuple[int, str, str], ...] = (),
) -> bytes:
    document = Document()
    if sections:
        for level, title, body in sections:
            document.add_heading(title, level=level)
            if body:
                document.add_paragraph(body)
    else:
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
